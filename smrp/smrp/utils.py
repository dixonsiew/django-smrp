from datetime import datetime, timezone
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from smrp.controllers.report.models import ColumnMap
import re, io


def get_date_str(v):
    o = ""
    
    # Check if v is a bson.DateTime (Python: datetime)
    if isinstance(v, datetime):
        o = get_str(v)
        try:
            iv = int(o)
        except (ValueError, TypeError):
            iv = 0
        t = datetime.fromtimestamp(iv / 1000, tz=timezone.utc)
        o = t.strftime("%Y-%m-%d")
        
    else:
        o = get_str(v)

    s = o
    if len(o) >= 10:
        i = s.find("/")
        if i > 0:
            try:
                s = s.split(" ")[0]
                g = datetime.strptime(s, "%d/%m/%Y")
                o = g.strftime("%Y-%m-%d")
            except ValueError:
                pass
            
        else:
            i = s.find("-")
            if i > 0:
                try:
                    s = s.split(" ")[0]
                    g = datetime.strptime(s, "%d-%m-%Y")
                    o = g.strftime("%Y-%m-%d")
                except ValueError:
                    pass
            
        s = o.split(" ")[0]
        
    else:
        i = s.find("/")
        if i > 0:
            try:
                s = s.split(" ")[0]
                g = datetime.strptime(s, "%d/%m/%Y")
                o = g.strftime("%Y-%m-%d")
            except ValueError:
                pass
            
        else:
            i = s.find("-")
            if i > 0:
                try:
                    s = s.split(" ")[0]
                    g = datetime.strptime(s, "%d-%m-%Y")
                    o = g.strftime("%Y-%m-%d")
                except ValueError:
                    pass
                
        s = o.split(" ")[0]

    return s

def set_value(x, ofield, src_field):
    v = x[src_field]
    if ofield in x:
        s = x[ofield]
        if s == "N/A":
            x[ofield] = v
            
    else:
        x[ofield] = v

    if x[ofield] == "undefined":
        x[ofield] = "N/A"

def get_str(a):
    return str(a)

def get_number(s):
    try:
        return int(s)
    except:
        return 0


def get_num(s):
    r = re.sub(r"[^\d.]*", "", s)
    try:
        return float(r)
    except:
        return 0.0
    
def process_doc(lx: list):
    ls = []
    na = "N/A"
    for x in lx:
        if "ADMISSION_DATE" in x:
            x["ADMISSION_DATE"] = get_date_str(x["ADMISSION_DATE"])

        if "DISCHARGE_DATE" in x:
            x["DISCHARGE_DATE"] = get_date_str(x["DISCHARGE_DATE"])

        if "DEATH_DATE" in x:
            x["DEATH_DATE"] = get_date_str(x["DEATH_DATE"])

        if "DELIVERY_DATE" in x:
            x["DELIVERY_DATE"] = get_date_str(x["DELIVERY_DATE"])

        if "PATIENT_NOK_NAME" in x:
            s = x["PATIENT_NOK_NAME"]
            if na == s:
                x["NOK_STREET1"] = na
                x["NOK_STREET2"] = na
                x["NOK_CITYCODE"] = na
                x["NOK_POSTCODE"] = na
                x["NOK_OCITY"] = na
                x["NOK_NATIONALITY"] = na
                
            else:
                set_value(x, "NOK_STREET1", "STREET1")
                set_value(x, "NOK_STREET2", "STREET2")
                set_value(x, "NOK_CITYCODE", "CITYCODE")
                set_value(x, "NOK_POSTCODE", "POSTCODE")
                set_value(x, "NOK_OCITY", "OCITY")
                set_value(x, "NOK_NATIONALITY", "NATIONALITY")
                
        else:
            x["NOK_STREET1"] = na
            x["NOK_STREET2"] = na
            x["NOK_CITYCODE"] = na
            x["NOK_POSTCODE"] = na
            x["NOK_OCITY"] = na
            x["NOK_NATIONALITY"] = na

        ls.append(x)

    return ls

def get_xlsx(colmaps: list[ColumnMap], lx: list):
    """
    colmaps: list of objects with .text and .field attributes (like report.ColumnMap)
    lx: list of dicts (like bson.M)
    Returns: io.BytesIO buffer
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    bold_font = Font(bold=True)
    coloffset = 6

    # Track column widths locally (openpyxl width is in characters)
    col_widths = {}

    # --- Header row ---
    for i, cx in enumerate(colmaps):
        j = i + 1
        cell = ws.cell(row=1, column=j, value=cx.text)
        cell.font = bold_font
        n = len(cx.text) if cx.text else 0
        col_letter = get_column_letter(j)
        col_widths[col_letter] = float(n + coloffset)
        ws.column_dimensions[col_letter].width = col_widths[col_letter]

    # --- Data rows ---
    k = 2
    for x in lx:
        for i, cx in enumerate(colmaps):
            field = cx.field
            j = i + 1
            s = ""
            if field in x:
                s = get_str(x[field])

            ws.cell(row=k, column=j, value=s)

            n = len(s) if s else 0
            col_letter = get_column_letter(j)
            m = col_widths.get(col_letter, 0.0)
            if float(n) > (m - float(coloffset)):
                new_width = float(n + coloffset)
                col_widths[col_letter] = new_width
                ws.column_dimensions[col_letter].width = new_width

        k += 1

    # --- Write to buffer ---
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf